"""One report snapshot shared by the GUI, local storage and every export format."""
from __future__ import annotations

import csv
from datetime import datetime
import json
from pathlib import Path
from typing import Any, Mapping
from xml.sax.saxutils import escape

from prediction_core import APP_VERSION, FEATURES, Inputs, PredictionModel

RESEARCH_NOTICE = {
    "en": "Research use only. This is not a diagnosis or a calibrated individual disease probability. Clinical assessment and confirmatory investigations remain necessary. Performance is internal to a selected paediatric THA/IDA cohort; healthy participants and other causes of anaemia were not evaluated.",
    "zh": "仅供研究。结果不是诊断，也不是经校准的个人患病概率；仍需临床评估和确诊检查。所列效能来自特定儿童地中海贫血/缺铁性贫血队列的内部评估，未评价健康人或其他贫血病因。",
}


def create_record(model: PredictionModel, values: Mapping[str, Any], *, hb_unit: str = "g/dL",
                  patient: Mapping[str, str] | None = None) -> dict[str, Any]:
    inputs = Inputs.from_mapping(values, hb_unit=hb_unit)
    prediction = model.predict(inputs)
    return {
        "record_schema_version": "1.0",
        "application_version": APP_VERSION,
        "created_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "patient": {str(key): str(value) for key, value in (patient or {}).items()},
        "entered_inputs": {key: float(values[key]) for key in FEATURES},
        "entered_Hb_unit": hb_unit,
        "prediction": prediction.to_dict(),
        "equation": model.equation(),
        "performance": model.document.get("performance", {}),
        "research_notice": dict(RESEARCH_NOTICE),
    }


def _metric(values: Mapping[str, Any], *names: str) -> float | None:
    for name in names:
        value = values.get(name)
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return float(value)
    return None


def benchmark_text(values: Mapping[str, Any], lang: str) -> str:
    auc = _metric(values, "AUC_score", "AUC")
    sensitivity = _metric(values, "Sensitivity", "sensitivity")
    specificity = _metric(values, "Specificity", "specificity")
    text = []
    if auc is not None:
        text.append(f"{'排序分数AUC' if lang == 'zh' else 'Score AUC'} {auc:.3f}")
    if sensitivity is not None:
        text.append(f"{'敏感度' if lang == 'zh' else 'Sensitivity'} {sensitivity:.1%}")
    if specificity is not None:
        text.append(f"{'特异度' if lang == 'zh' else 'Specificity'} {specificity:.1%}")
    return "; ".join(text)


def report_text(record: Mapping[str, Any], lang: str = "en") -> str:
    zh = lang == "zh"
    p = record["prediction"]
    lines = ["THA/IDA 研究计算报告" if zh else "THA/IDA research calculation report",
             f"{'模型版本' if zh else 'Model version'}: {p['model_version']} ({p['model_id']})",
             f"{'程序版本' if zh else 'Application version'}: {record['application_version']}",
             f"{'生成时间' if zh else 'Created'}: {record['created_at']}"]
    labels = {"name": ("姓名", "Name"), "id": ("患者ID", "Patient ID"), "age": ("年龄（岁）", "Age (years)"),
              "gender": ("性别", "Sex"), "hometown": ("籍贯", "Hometown"), "test_date": ("检测日期", "Test date")}
    for key, value in record["patient"].items():
        if value:
            label = labels.get(key, (key, key))[0 if zh else 1]
            lines.append(f"{label}: {value}")
    positive = "阳性" if zh else "positive"
    negative = "阴性" if zh else "negative"
    lines.extend(["", f"{'联合OR规则' if zh else 'Combined OR rule'}: {positive if p['combined_positive'] else negative}",
                  f"{'逻辑分数' if zh else 'Logistic score'} q = {p['probability']:.6f}; "
                  f"{'阳性指数数' if zh else 'Positive index votes'} = {p['positive_votes']}/{len(p['rules'])}",
                  "仅供研究；不可替代诊断。" if zh else "Research only; this is not a diagnosis."])
    lines.extend(["", "输入（统一单位）" if zh else "Inputs (canonical units)"])
    for key in FEATURES:
        lines.append(f"{key}: {p['inputs'][key]:.12g} {p['input_units'][key]}")
    lines.append(f"{'录入Hb' if zh else 'Entered Hb'}: {record['entered_inputs']['Hb']:.12g} {record['entered_Hb_unit']}")
    lines.extend(["", "五指数投票" if zh else "Five-index voting"])
    for rule in p["rules"]:
        decision = positive if rule["positive"] else negative
        lines.append(f"{rule['name']}: {rule['value']:.12g}; {rule['operator']} {rule['threshold']:.17g}; {decision}")
        lines.append(f"  {rule['expression']}")
        benchmark = benchmark_text(rule.get("full_cohort", {}), lang)
        if benchmark:
            lines.append(f"  {'开发队列表观效能' if zh else 'Development-cohort apparent performance'}: {benchmark}")
    lines.extend(["", record["equation"], f"z = {p['logit']:.17g}", f"q = {p['probability']:.17g}",
                  f"q > {p['probability_threshold']:.17g}: {positive if p['student_positive'] else negative}",
                  f"V = {p['positive_votes']}/{len(p['rules'])}; V >= {p['required_positive_votes']}: {positive if p['voting_positive'] else negative}",
                  f"{'联合OR规则' if zh else 'Combined OR rule'}: {positive if p['combined_positive'] else negative}",
                  f"{'排序分数S（非个人风险）' if zh else 'Ranking score S (not individual risk)'} = {p['ranking_score']:.17g}"])
    performance = record.get("performance", {}).get("fixed_full_cohort", {})
    lines.extend(["", "开发队列表观效能（并非外部验证）" if zh else "Development-cohort apparent performance (not external validation)"])
    for name in ("DF", "Distill", "Combined"):
        if name in performance:
            lines.append(f"{name}: {benchmark_text(performance[name], lang)}")
    lines.extend(["", RESEARCH_NOTICE[lang], f"Model SHA256: {p['model_sha256']}"])
    return "\n".join(lines) + "\n"


def flat_record(record: Mapping[str, Any]) -> dict[str, Any]:
    p = record["prediction"]
    result = {"application_version": record["application_version"], "model_id": p["model_id"],
              "model_version": p["model_version"], "model_sha256": p["model_sha256"],
              "created_at": record["created_at"], "entered_Hb": record["entered_inputs"]["Hb"],
              "entered_Hb_unit": record["entered_Hb_unit"]}
    result.update({"patient_" + key: value for key, value in record["patient"].items()})
    result.update({key + "_" + p["input_units"][key]: value for key, value in p["inputs"].items()})
    for key in ("logit", "probability", "probability_threshold", "positive_votes", "required_positive_votes",
                "voting_positive", "student_positive", "combined_positive", "ranking_score"):
        result[key] = p[key]
    for rule in p["rules"]:
        for key in ("value", "operator", "threshold", "positive"):
            result[rule["expression_id"] + "_" + key] = rule[key]
    result["equation"] = record["equation"]
    result["research_notice"] = RESEARCH_NOTICE["en"]
    return result


def _spreadsheet_value(value: Any) -> Any:
    # Preserve user metadata as literal text in CSV and Excel.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def export_record(record: Mapping[str, Any], path: str | Path, lang: str = "en") -> None:
    path = Path(path)
    extension = path.suffix.lower()
    if extension == ".json":
        path.write_text(json.dumps(record, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    elif extension == ".txt":
        path.write_text(report_text(record, lang), encoding="utf-8")
    elif extension == ".csv":
        row = flat_record(record)
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=list(row))
            writer.writeheader()
            writer.writerow({key: _spreadsheet_value(value) for key, value in row.items()})
    elif extension == ".xlsx":
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment
        except ImportError as exc:
            raise RuntimeError("Excel export requires openpyxl; install requirements.txt") from exc
        book = Workbook()
        summary = book.active
        summary.title = "Research report"
        summary.append(["Report"])
        for line in report_text(record, lang).splitlines():
            summary.append([_spreadsheet_value(line)])
        summary.column_dimensions["A"].width = 120
        for row in summary:
            row[0].alignment = Alignment(wrap_text=True, vertical="top")
        summary["A1"].font = Font(bold=True)
        data = book.create_sheet("Structured results")
        row = flat_record(record)
        data.append(list(row))
        data.append([_spreadsheet_value(value) for value in row.values()])
        data.freeze_panes = "A2"
        snapshot = book.create_sheet("JSON snapshot")
        # Split to avoid Excel's 32,767-character cell limit.
        for line in json.dumps(record, ensure_ascii=False, indent=2, allow_nan=False).splitlines():
            snapshot.append([_spreadsheet_value(line)])
        book.save(path)
    elif extension == ".pdf":
        try:
            from reportlab.lib.styles import ParagraphStyle
            from reportlab.lib.pagesizes import A4
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        except ImportError as exc:
            raise RuntimeError("PDF export requires reportlab; install requirements.txt") from exc
        # Chinese patient metadata can occur in either language's report.
        font = "STSong-Light"
        if font not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(UnicodeCIDFont(font))
        style = ParagraphStyle("ResearchReport", fontName=font, fontSize=9, leading=13,
                               wordWrap="CJK" if lang == "zh" else None, splitLongWords=True)
        content = [Paragraph(escape(line), style) if line else Spacer(1, 7)
                   for line in report_text(record, lang).splitlines()]
        SimpleDocTemplate(str(path), pagesize=A4, leftMargin=36, rightMargin=36,
                          topMargin=36, bottomMargin=36, title="THA/IDA research calculation report").build(content)
    else:
        raise ValueError("Supported formats: .pdf, .xlsx, .csv, .txt, .json")
